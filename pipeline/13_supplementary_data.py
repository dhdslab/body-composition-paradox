"""Step 13. Supplementary Data 1 (aggregated results of every analysis) and the Source Data file (data behind each figure panel).

Outputs:
    results/Supplementary_Data_1.xlsx
    results/Source_Data.xlsx
"""
import json

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from common import DISEASE_NAMES as NAMES, MODELS, RESULTS, SOURCE_DATA, TABLES, pd

MODEL = {'A_anthropometry': 'A: anthropometry', 'B_inferred_body': 'B: A + estimated composition', 'C_laboratory': 'C: A + laboratory', 'D_laboratory_body': 'D: C + estimated composition',
         'age_sex': 'E: age + sex', 'measured_body': 'F: age + sex + measured DXA', 'anthropometry': 'G: anthropometry', 'anthropometry_measured': 'H: G + measured DXA'}


def eng(df, col='disease'):
    df = df.copy()
    if 'label' in df: df = df.drop(columns='label')
    if col in df: df.insert(df.columns.get_loc(col) + 1, 'endpoint', df[col].map(NAMES))
    if 'model' in df and df.model.isin(MODEL.keys()).any(): df['model'] = df.model.map(lambda m: MODEL.get(m, m))
    return df


def manifest():
    rows = []
    for name, era in [('body_model_manifest.json', 'Historical body composition (2008–2010)'), ('contemporary_manifest.json', 'Contemporary body composition (2024)')]:
        for m in json.loads((MODELS / name).read_text()):
            rows.append(dict(model_set=era, target=m['target'], tier=m['tier'], family=m['family'], features='; '.join(m['features']), validation=m.get('validation', '2024 temporal test of frozen model')))
    dv = pd.read_csv(TABLES / 'disease_validation.csv')
    for _, r in dv[dv.period.eq('2019_2021')].iterrows():
        rows.append(dict(model_set='Disease models A–D (2011–2018)', target=NAMES[r.disease], tier=MODEL[r.model], family=r.family, features='see Methods', validation='2019–2021 temporal test'))
    for m in json.loads((MODELS / 'measured_disease_manifest.json').read_text()):
        rows.append(dict(model_set=f"Measured-DXA disease model F ({m['era']})", target=NAMES[m['disease']], tier='F', family='spline logistic', features='; '.join(m['features']), validation=m['validation']))
    return pd.DataFrame(rows)


sheets = [
    ('01_cohort_flow', 'Participants by survey year and eligibility step', pd.read_csv(TABLES / 'cohort_flow.csv')),
    ('02_predictor_missing', 'Proportion missing for each predictor by survey year (eligible adults)', pd.read_csv(TABLES / 'predictor_missingness.csv')),
    ('03_table1', 'Characteristics of the analytic populations (Table 1)', pd.read_csv(TABLES / 'Table_1_characteristics.csv')),
    ('04_body_family_tuning', 'Body composition model family selection in 2010 (2008–2009 fit)', pd.read_csv(TABLES / 'body_tuning.csv')),
    ('05_body_valid_hist', 'Historical body composition models: 2010 age-matched tuning and 2024 test, all tiers, with 95% CIs', pd.read_csv(TABLES / 'body_validation.csv')),
    ('06_body_valid_2024cv', '2024 models, fivefold PSU cross-validation, all tiers, with 95% CIs', pd.read_csv(TABLES / 'contemporary_body_validation.csv')),
    ('07_body_sub_hist', 'Historical models in 2024 by sex and age group, all tiers', pd.read_csv(TABLES / 'body_subgroups.csv')),
    ('08_body_sub_2024cv', '2024 models by sex and age group, all tiers (new)', pd.read_csv(TABLES / 'contemporary_body_subgroups.csv')),
    ('09_body_sensitivity', 'Historical models against complete-laboratory DXA and BIA references (all ages for 2022–2023 BIA)', pd.read_csv(TABLES / 'body_sensitivity.csv')),
    ('10_bia_dxa_2024', 'Paired 2024 BIA minus DXA agreement', pd.read_csv(TABLES / 'BIA_DXA_agreement_2024.csv')),
    ('11_reference_triang', 'Historical estimates versus 2010 DXA, 2022–2023 BIA and paired 2024 BIA/DXA, age ≥40 (new)', pd.read_csv(TABLES / 'bias_triangulation.csv')),
    ('12_era_difference', 'Measured DXA, 2024 minus 2008–2010, unadjusted and anthropometry-adjusted, by group, as released (new)', pd.read_csv(TABLES / 'era_difference_measured_dxa.csv')),
    ('12b_era_diff_recalibrated', 'As sheet 12 after recalibrating 2008–2010 to the NHANES calibration (lean soft tissue x 0.946) (new)', pd.read_csv(TABLES / 'era_difference_recalibrated.csv')),
    ('13_era_by_bmi', 'Weighted mean measured DXA by BMI category, period and sex, with 95% stratified PSU bootstrap intervals (new)', pd.read_csv(TABLES / 'era_binned_by_bmi.csv')),
    ('14_low_almi', 'Proportion below AWGS 2019 DXA low-muscle-mass thresholds by period, sex and age, as released (new)', pd.read_csv(TABLES / 'low_almi_by_era.csv')),
    ('14b_low_almi_recalibrated', 'As sheet 14 after recalibrating 2008–2010 appendicular lean mass to the NHANES calibration (x 0.946) (new)', pd.read_csv(TABLES / 'low_almi_recalibrated.csv')),
    ('15_disease_catalogue', 'All 49 current-disease items, counts and disposition', eng(pd.read_csv(TABLES / 'disease_catalogue.csv'), 'variable')),
    ('16_disease_family_tune', 'Disease model family selection by 2017–2018 log loss', eng(pd.read_csv(TABLES / 'disease_tuning.csv'))),
    ('17_disease_validation', 'Disease models A–D, 2019–2021 and 2022–2024, all metrics with 95% CIs', eng(pd.read_csv(TABLES / 'disease_validation.csv'))),
    ('17b_disease_est_only', 'Disease models with only the four estimated body composition values (laboratory-assisted or anthropometric estimates) and with age and sex only, 2019–2021, all metrics with 95% CIs (new)',
     eng(pd.read_csv(TABLES / 'disease_estimates_only.csv').replace({'model': {'estimates_only_laboratory': 'Estimated composition only (laboratory-assisted estimates)',
                                                                            'estimates_only_anthropometry': 'Estimated composition only (anthropometric estimates)', 'age_sex': 'Age and sex only'}}))),
    ('18_disease_paired', 'Paired differences between disease models with Benjamini–Hochberg q values', eng(pd.read_csv(TABLES / 'disease_paired_comparisons.csv'))),
    ('18b_disease_same_family', 'Paired AUROC differences, model B versus A and model D versus C, with model B or D refitted using the family selected for its comparator when the two differed (post hoc) (new)', eng(pd.read_csv(TABLES / 'disease_same_family.csv'))),
    ('19_disease_subgroups', 'Disease models A–D by sex and age group, 2019–2021', eng(pd.read_csv(TABLES / 'disease_subgroups.csv'))),
    ('20_calib_25_endpoints', 'Weighted calibration deciles, models B and C, 2019–2021', eng(pd.read_csv(TABLES / 'calibration_deciles.csv'))),
    ('21_calib_recurring', 'Weighted calibration deciles, models A–D, 2019–2021 and 2022–2024 (new)', pd.read_csv(TABLES / 'calibration_deciles_recurring.csv')),
    ('22_diagnostic_lab_omit', 'Model C after omitting diagnosis-related biomarkers', eng(pd.read_csv(TABLES / 'disease_diagnostic_lab_sensitivity.csv'))),
    ('23_measured_catalogue', 'Eligibility for measured-DXA disease models', eng(pd.read_csv(TABLES / 'measured_body_disease_catalogue.csv'))),
    ('24_measured_validation', 'Measured-DXA disease models E–H: 2010, 2024 transport and 2024 cross-validation', eng(pd.read_csv(TABLES / 'measured_body_disease_validation.csv'))),
    ('25_measured_paired', 'Paired AUROC and Brier differences for measured-DXA models, 2010', eng(pd.read_csv(TABLES / 'measured_body_disease_comparisons.csv'))),
    ('26_measured_calib', 'Calibration deciles for model F in 2024 (new)', pd.read_csv(TABLES / 'calibration_deciles_measured.csv')),
    ('27_model_specs', 'Selected model families and predictors for all stored models', manifest())]

GLOSS = [('rmse / mae / r2 / bias', 'Survey-weighted root mean squared error, mean absolute error, R² and mean (predicted − reference); percentage points for body_fat_pct, kg otherwise'),
         ('loa_low / loa_high', 'Bias ± 1.96 survey-weighted standard deviations of the differences'),
         ('*_lo / *_hi', '95% percentile interval from 500 Rao–Wu rescaled bootstrap replicates of primary sampling units within year-specific strata, predictions held fixed'),
         ('auc / ap / brier / logloss', 'Survey-weighted AUROC, average precision, Brier score and log loss'),
         ('prevalence / predicted_prevalence', 'Survey-weighted observed proportion of cases and mean predicted probability'),
         ('calibration_intercept / calibration_slope', 'Calibration-in-the-large with slope fixed at 1; slope estimated jointly with a free intercept'),
         ('q_value', 'Benjamini–Hochberg adjusted P value within each comparison and metric'),
         ('tier', 'minimal = age, sex, BMI; anthropometry = + height, weight, waist; laboratory = + 12 blood measurements'),
         ('period', '2010_tuning_age40, 2024_external, 2019_2021, 2022_2024, 2010_temporal, 2024_transport, 2024_internal_psu_cv'),
         ('n / train_n / test_n / events', 'Unweighted counts')]
with pd.ExcelWriter(RESULTS / 'Supplementary_Data_1.xlsx', engine='openpyxl') as xw:
    readme = pd.DataFrame({'Supplementary Data 1': ['Body Composition Paradox: Shifting Reference Standards and Limited Added Value of Estimates in Disease Prediction',
                                                     'Aggregated results for all analyses in the manuscript and Supplementary Information. No participant-level data are included.',
                                                     'Source: Korea National Health and Nutrition Examination Survey (KNHANES) 2008–2024 public-use data, Korea Disease Control and Prevention Agency.', '']})
    readme.to_excel(xw, sheet_name='README', index=False, startrow=0)
    pd.DataFrame([dict(sheet=s, contents=d, rows=len(df)) for s, d, df in sheets]).to_excel(xw, sheet_name='README', index=False, startrow=6)
    pd.DataFrame(GLOSS, columns=['column(s)', 'definition']).to_excel(xw, sheet_name='README', index=False, startrow=9 + len(sheets))
    for s, d, df in sheets: df.to_excel(xw, sheet_name=s, index=False)
    for ws in xw.book.worksheets:
        ws.freeze_panes = 'A2' if ws.title != 'README' else None
        for c in ws[1]: c.font = Font(bold=True)
        for i, col in enumerate(ws.columns, 1):
            width = min(60, max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]) + 2)
            ws.column_dimensions[get_column_letter(i)].width = max(10, width)
    ws = xw.book['README']; ws['A1'].font = Font(bold=True, size=13); ws.column_dimensions['B'].width = 100

order = ['Fig1', 'Fig1_by_year', 'Fig2', 'Fig3a', 'Fig3b_c', 'Fig3d', 'Fig3e', 'Fig4a', 'Fig4b_c', 'Fig5a_c', 'Fig5d', 'Fig6a', 'Fig6b', 'Fig6c', 'Fig6d',
         'SupplementaryFig1', 'SupplementaryFig2', 'SupplementaryFig3', 'SupplementaryFig4']
with pd.ExcelWriter(RESULTS / 'Source_Data.xlsx', engine='openpyxl') as xw:
    pd.DataFrame({'Source Data': ['Source data for Figs. 1–6 and Supplementary Figs. 1–4 of “Body Composition Paradox: Shifting Reference Standards and Limited Added Value of Estimates in Disease Prediction”.',
                                  'Each sheet contains the aggregated values plotted in the named panel. Supplementary Fig. 1 is a density plot of participant-level differences; only the plotted summary lines are provided, because participant-level KNHANES data cannot be redistributed.',
                                  'All estimates are survey-weighted unless labelled n.']}).to_excel(xw, sheet_name='README', index=False)
    for s in order:
        df = pd.read_csv(SOURCE_DATA / f'{s}.csv')
        if 'disease' in df and 'endpoint' not in df: df = eng(df)
        df.to_excel(xw, sheet_name=s, index=False)
    for ws in xw.book.worksheets:
        for c in ws[1]: c.font = Font(bold=True)
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(70, max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]) + 2))
print('Supplementary Data 1 sheets:', len(sheets), '| Source Data sheets:', len(order))
